"""REV25: TAGs/Produtos — banco próprio, importação pré-validada e exportação."""
from datetime import datetime
from functools import wraps
from io import BytesIO
from PIL import Image, UnidentifiedImageError
from sqlalchemy import text
import re, unicodedata, json
from flask import request, session, jsonify, abort, redirect, url_for, render_template, send_file
from sqlalchemy import func
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

def norm(s):
    s=unicodedata.normalize('NFKD',str(s or '')).encode('ascii','ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]+','_',s).strip('_')
FIELDS=['serial','patrimonio','produto','lote','empresa','localidade','finalidade','status','nf_entrada','nf_saida','caixa','observacao']
ALIASES={
 'serial':['N_SERIE','N_DE_SERIE','NUMERO_DE_SERIE','NUMERO_SERIE','SERIE','SERIAL','TAG','N_SERIE_TAG','N_SERIE_BF80','N_S','NSERIE'],
 'patrimonio':['PATRIMONIO','N_PATRIMONIO'], 'produto':['PRODUTO','MODELO','EQUIPAMENTO'],
 'lote':['LOTE','PRODUCAO','N_PRODUCAO','FABRICACAO'],
 'empresa':['EMPRESA','PROPRIETARIO','EMPRESA_PROPRIETARIA'],
 'localidade':['LOCAL_PROJETO','LOCALIDADE','LOCAL','PROJETO','DESTINO','ENDERECAMENTO'],
 'finalidade':['FINALIDADE','TIPO_OPERACAO'],
 'status':['STATUS','STATUS_OPERACIONAL','SITUACAO'],
 'nf_entrada':['NF_ENTRADA','NOTA_FISCAL_ENTRADA','NF_DE_ENTRADA'],
 'nf_saida':['NF_SAIDA','NOTA_FISCAL_SAIDA','NF_DE_SAIDA'],
 'caixa':['CAIXA','N_CAIXA'], 'observacao':['OBSERVACAO','OBSERVACOES','OBS']}
ALIASES={k:set([norm(k)]+[norm(x) for x in v]) for k,v in ALIASES.items()}

def install_tags(app,db,can):
 class Tag(db.Model):
    __tablename__='v87_product_tags'
    id=db.Column(db.Integer,primary_key=True)
    serial=db.Column(db.String(160),nullable=False)
    produto=db.Column(db.String(160),nullable=False,default='BF80')
    patrimonio=db.Column(db.String(160),default='')
    lote=db.Column(db.String(120),default='')
    empresa=db.Column(db.String(160),default='')
    localidade=db.Column(db.String(200),default='')
    finalidade=db.Column(db.String(120),default='')
    status=db.Column(db.String(120),default='')
    nf_entrada=db.Column(db.String(120),default='')
    nf_saida=db.Column(db.String(120),default='')
    caixa=db.Column(db.String(120),default='')
    observacao=db.Column(db.Text,default='')
    created_at=db.Column(db.DateTime,default=datetime.utcnow)
    updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)
    __table_args__=(db.UniqueConstraint('produto','serial',name='uq_v87_tags_product_serial'),)
 class Product(db.Model):
    __tablename__='v87_tag_products'
    id=db.Column(db.Integer,primary_key=True)
    modelo=db.Column(db.String(160),unique=True,nullable=False)
    nome=db.Column(db.String(200),nullable=False)
    fabricante=db.Column(db.String(160),default='')
    categoria=db.Column(db.String(160),default='')
    descricao=db.Column(db.Text,default='')
    foto=db.Column(db.LargeBinary,nullable=True)
    ativo=db.Column(db.Boolean,nullable=False,default=True)
    created_at=db.Column(db.DateTime,default=datetime.utcnow)
 class TagPhoto(db.Model):
    __tablename__='v87_product_tag_photos'
    id=db.Column(db.Integer,primary_key=True)
    tag_id=db.Column(db.Integer,nullable=False,index=True)
    image=db.Column(db.LargeBinary,nullable=False)
    mime=db.Column(db.String(32),nullable=False)
    description=db.Column(db.String(250),default='')
    user_id=db.Column(db.Integer,nullable=False)
    created_at=db.Column(db.DateTime,default=datetime.utcnow)
 class Audit(db.Model):
    __tablename__='v87_product_tag_audit'
    id=db.Column(db.Integer,primary_key=True)
    tag_id=db.Column(db.Integer,nullable=False)
    user_id=db.Column(db.Integer,nullable=False)
    action=db.Column(db.String(30),nullable=False)
    before_json=db.Column(db.Text)
    after_json=db.Column(db.Text)
    created_at=db.Column(db.DateTime,default=datetime.utcnow)
 class ImportIssue(db.Model):
    __tablename__='v87_product_tag_import_issues'
    id=db.Column(db.Integer,primary_key=True)
    produto=db.Column(db.String(160),nullable=False)
    serial=db.Column(db.String(160),nullable=False)
    lines=db.Column(db.String(300),default='')
    detail=db.Column(db.Text,default='')
    file_name=db.Column(db.String(255),default='')
    user_id=db.Column(db.Integer,nullable=False)
    created_at=db.Column(db.DateTime,default=datetime.utcnow)
    resolved_at=db.Column(db.DateTime,nullable=True)
 initialized=False
 def guard(permission="engineering.tags.view"):
    def outer(fn):
     @wraps(fn)
     def inner(*a,**kw):
      nonlocal initialized
      if not session.get('user_id'):
       return (jsonify(ok=False,error='Autenticação necessária'),401) if request.path.startswith('/api/') else redirect(url_for('login'))
      if not can(permission): abort(403)
      if not initialized:
       Tag.__table__.create(bind=db.engine,checkfirst=True)
       Product.__table__.create(bind=db.engine,checkfirst=True)
       with db.engine.begin() as conn:
        conn.execute(text('ALTER TABLE v87_tag_products ADD COLUMN IF NOT EXISTS foto BYTEA'))
        conn.execute(text('ALTER TABLE v87_tag_products ADD COLUMN IF NOT EXISTS ativo BOOLEAN NOT NULL DEFAULT TRUE'))
       TagPhoto.__table__.create(bind=db.engine,checkfirst=True)
       Audit.__table__.create(bind=db.engine,checkfirst=True)
       ImportIssue.__table__.create(bind=db.engine,checkfirst=True)
       initialized=True
      return fn(*a,**kw)
     return inner
    return outer
 def asdict(t):return {k:str(getattr(t,k) or '') for k in FIELDS}|{'id':t.id,'updated_at':t.updated_at.isoformat() if t.updated_at else ''}
 def query_filtered():
    q=Tag.query
    for field in ['produto','empresa','localidade','finalidade','status','lote']:
     val=request.args.get(field,'').strip()
     if val:q=q.filter(getattr(Tag,field)==val)
    for field,param in [('nf_entrada','nf_entrada_filtro'),('nf_saida','nf_saida_filtro')]:
     val=request.args.get(param,'').strip()
     if val:q=q.filter(getattr(Tag,field).ilike('%'+val+'%'))
    situation=request.args.get('situacao_nf','').strip()
    has_entry=db.and_(Tag.nf_entrada.isnot(None),db.func.trim(Tag.nf_entrada)!='')
    has_exit=db.and_(Tag.nf_saida.isnot(None),db.func.trim(Tag.nf_saida)!='')
    if situation=='entrada':q=q.filter(has_entry)
    elif situation=='saida':q=q.filter(has_exit)
    elif situation=='ambas':q=q.filter(has_entry,has_exit)
    elif situation=='nenhuma':q=q.filter(db.not_(has_entry),db.not_(has_exit))
    search=request.args.get('q','').strip()
    if search:q=q.filter(db.or_(Tag.serial.ilike('%'+search+'%'),Tag.patrimonio.ilike('%'+search+'%'),Tag.localidade.ilike('%'+search+'%')))
    return q
 @app.get('/api/engenharia/tags/produtos')
 @guard()
 def tags_products():
    registered=Product.query.order_by(Product.nome).all()
    known={x.modelo.casefold() for x in registered}
    products=[dict(modelo=x.modelo,nome=x.nome,fabricante=x.fabricante or '',categoria=x.categoria or '',descricao=x.descricao or '',tem_foto=bool(x.foto),ativo=x.ativo) for x in registered]
    for (modelo,) in db.session.query(Tag.produto).distinct().order_by(Tag.produto).all():
     if modelo and modelo.casefold() not in known:
      products.append(dict(modelo=modelo,nome='Biometria Facial' if modelo.upper()=='BF80' else modelo,fabricante='Autopass' if modelo.upper()=='BF80' else '',categoria='Biometria Facial' if modelo.upper()=='BF80' else '',descricao='',tem_foto=False,ativo=True))
    if not any(x['modelo'].upper()=='BF80' for x in products):
     products.append(dict(modelo='BF80',nome='Biometria Facial',fabricante='Autopass',categoria='Biometria Facial',descricao='Equipamento de reconhecimento facial',tem_foto=False,ativo=True))
    return jsonify(ok=True,items=sorted(products,key=lambda x:x['modelo'].casefold()))
 @app.post('/api/engenharia/tags/produtos')
 @guard('engineering.tags.import')
 def tags_create_product():
    data=request.form
    modelo=str(data.get('modelo') or '').strip().upper()
    nome=str(data.get('nome') or '').strip()
    if not modelo or not nome or len(modelo)>160 or len(nome)>200:
     return jsonify(ok=False,error='Modelo e nome são obrigatórios (máximo 160/200 caracteres).'),400
    if Product.query.filter(func.lower(Product.modelo)==modelo.lower()).first() or db.session.query(Tag.id).filter(func.lower(Tag.produto)==modelo.lower()).first():
     return jsonify(ok=False,error='Modelo já existente. Selecione-o na lista.'),409
    uploaded=request.files.get('foto')
    if not uploaded or not uploaded.filename:
     return jsonify(ok=False,error='Fotografia obrigatória para o novo produto.'),400
    raw=uploaded.read(3*1024*1024+1)
    if len(raw)>3*1024*1024:
     return jsonify(ok=False,error='A fotografia excede 3 MB.'),413
    try:
     with Image.open(BytesIO(raw)) as im:
      if im.format not in ('JPEG','PNG','WEBP'):
       raise ValueError('Formato inválido')
      im.verify()
     with Image.open(BytesIO(raw)) as im:
      im.thumbnail((1200,1200))
      if im.mode not in ('RGB','L'):im=im.convert('RGB')
      output=BytesIO();im.save(output,format='JPEG',quality=85,optimize=True)
      photo=output.getvalue()
    except (UnidentifiedImageError,ValueError,OSError,Image.DecompressionBombError):
     return jsonify(ok=False,error='Envie uma imagem JPG, PNG ou WebP válida.'),400
    item=Product(modelo=modelo,nome=nome,fabricante=str(data.get('fabricante') or '').strip()[:160],categoria=str(data.get('categoria') or '').strip()[:160],descricao=str(data.get('descricao') or '').strip()[:2000],foto=photo)
    db.session.add(item)
    db.session.commit()
    return jsonify(ok=True,produto=dict(modelo=item.modelo,nome=item.nome)),201
 def product_photo_upload(uploaded):
    if not uploaded or not uploaded.filename:return None, 'Selecione uma fotografia.'
    raw=uploaded.read(3*1024*1024+1)
    if len(raw)>3*1024*1024:return None, 'A fotografia excede 3 MB.'
    try:
     with Image.open(BytesIO(raw)) as im:
      if im.format not in ('JPEG','PNG','WEBP'):raise ValueError('Formato inválido')
      im.verify()
     with Image.open(BytesIO(raw)) as im:
      im.thumbnail((1200,1200))
      if im.mode not in ('RGB','L'):im=im.convert('RGB')
      output=BytesIO();im.save(output,format='JPEG',quality=85,optimize=True)
      return output.getvalue(),None
    except (UnidentifiedImageError,ValueError,OSError,Image.DecompressionBombError):
     return None,'Envie uma imagem JPG, PNG ou WebP válida.'
 def product_audit(action,model,before,after):
    db.session.add(Audit(tag_id=0,user_id=int(session.get('user_id') or 0),action=action,before_json=json.dumps({'modelo':model,**before},ensure_ascii=False),after_json=json.dumps({'modelo':model,**after},ensure_ascii=False)))
 @app.patch('/api/engenharia/tags/produtos/<path:modelo>')
 @guard('engineering.tags.edit')
 def tags_edit_product(modelo):
    item=Product.query.filter(func.lower(Product.modelo)==modelo.lower()).first()
    if not item:
     # Produtos legados aparecem no seletor mesmo sem cadastro de metadados.
     if not db.session.query(Tag.id).filter(func.lower(Tag.produto)==modelo.lower()).first() and modelo.upper()!='BF80':
      return jsonify(ok=False,error='Produto não encontrado.'),404
     item=Product(modelo=modelo.upper(),nome='Biometria Facial' if modelo.upper()=='BF80' else modelo.upper(),fabricante='Autopass' if modelo.upper()=='BF80' else '',categoria='Biometria Facial' if modelo.upper()=='BF80' else '')
     db.session.add(item)
    data=request.form
    before={k:getattr(item,k) for k in ('nome','fabricante','categoria','descricao','ativo')}
    for k,maxlen in (('nome',200),('fabricante',160),('categoria',160),('descricao',2000)):
     if k in data:
      value=str(data.get(k) or '').strip()[:maxlen]
      if k=='nome' and not value:return jsonify(ok=False,error='Nome do produto obrigatório.'),400
      setattr(item,k,value)
    uploaded=request.files.get('foto')
    if uploaded and uploaded.filename:
     photo,error=product_photo_upload(uploaded)
     if error:return jsonify(ok=False,error=error),400
     item.foto=photo
    elif data.get('remover_foto')=='sim':item.foto=None
    after={k:getattr(item,k) for k in ('nome','fabricante','categoria','descricao','ativo')}
    after['foto_alterada']=bool(uploaded and uploaded.filename) or data.get('remover_foto')=='sim'
    product_audit('PRODUCT_EDIT',item.modelo,before,after)
    db.session.commit()
    return jsonify(ok=True,produto={'modelo':item.modelo,'nome':item.nome,'tem_foto':bool(item.foto),'ativo':item.ativo})
 @app.delete('/api/engenharia/tags/produtos/<path:modelo>')
 @guard('engineering.tags.edit')
 def tags_delete_product(modelo):
    item=Product.query.filter(func.lower(Product.modelo)==modelo.lower()).first()
    count=db.session.query(func.count(Tag.id)).filter(func.lower(Tag.produto)==modelo.lower()).scalar() or 0
    if not item and not count:return jsonify(ok=False,error='Produto não encontrado.'),404
    if request.args.get('confirmar')!='sim':return jsonify(ok=False,error='Confirmação obrigatória.'),400
    if count:
     if not item:
      item=Product(modelo=modelo.upper(),nome='Biometria Facial' if modelo.upper()=='BF80' else modelo.upper())
      db.session.add(item)
     if not item.ativo:return jsonify(ok=False,error='Produto já está inativo.'),409
     item.ativo=False
     product_audit('PRODUCT_DISABLE',item.modelo,{'ativo':True,'tags':count},{'ativo':False,'tags':count})
     db.session.commit()
     return jsonify(ok=True,acao='inativado',tags=count)
    # Não excluir registros de divergência pendentes ou histórico do produto.
    issues=db.session.query(func.count(ImportIssue.id)).filter(func.lower(ImportIssue.produto)==modelo.lower()).scalar() or 0
    if issues:return jsonify(ok=False,error='Produto possui divergências históricas; inative-o em vez de excluir.'),409
    if modelo.upper()=='BF80':return jsonify(ok=False,error='BF80 é o produto padrão; não pode ser excluído.'),409
    product_audit('PRODUCT_DELETE',item.modelo,{'nome':item.nome,'ativo':item.ativo},{'excluido':True})
    db.session.delete(item)
    db.session.commit()
    return jsonify(ok=True,acao='excluido',tags=0)
 @app.post('/api/engenharia/tags/produtos/<path:modelo>/reativar')
 @guard('engineering.tags.edit')
 def tags_reactivate_product(modelo):
    item=Product.query.filter(func.lower(Product.modelo)==modelo.lower()).first()
    if not item:return jsonify(ok=False,error='Produto não encontrado.'),404
    item.ativo=True
    product_audit('PRODUCT_REACTIVATE',item.modelo,{'ativo':False},{'ativo':True})
    db.session.commit()
    return jsonify(ok=True)
 @app.get('/api/engenharia/tags/produtos/foto/<path:modelo>')
 @guard()
 def tags_product_photo(modelo):
    item=Product.query.filter(func.lower(Product.modelo)==modelo.lower()).first()
    if not item or not item.foto:abort(404)
    response=send_file(BytesIO(item.foto),mimetype='image/jpeg',max_age=3600)
    response.headers['X-Content-Type-Options']='nosniff'
    return response
 @app.get('/engenharia/tags')
 @guard()
 def tags_page():return render_template('v87_tags.html',tag_permissions={k:bool(can('engineering.tags.'+k)) for k in ('view','import','edit','bulk','issues','export','audit')})
 @app.get('/api/engenharia/tags')
 @guard()
 def tags_list():
    q=query_filtered(); total=q.count()
    sort=request.args.get('sort','serial');sort=sort if sort in FIELDS else 'serial'
    col=getattr(Tag,sort)
    q=q.order_by(col.desc() if request.args.get('dir')=='desc' else col.asc(),Tag.id.asc())
    page=max(1,int(request.args.get('page',1)));limit=min(200,max(1,int(request.args.get('limit',30))))
    items=q.offset((page-1)*limit).limit(limit).all()
    base=query_filtered().order_by(None)
    groups={}
    for key in ['produto','empresa','localidade','finalidade','status','lote']:
     col=getattr(Tag,key)
     result=base.with_entities(col,func.count(Tag.id)).group_by(col).all()
     groups[key]=[{'name':name or 'Não informado','count':count} for name,count in sorted(result,key=lambda x:(-x[1],x[0] or ''))]
    from sqlalchemy import case
    summary_row=base.with_entities(
     func.sum(case((func.upper(Tag.finalidade).like('%LOC%'),1),else_=0)),
     func.sum(case((func.upper(Tag.finalidade).like('%VEND%'),1),else_=0)),
     func.sum(case((func.upper(Tag.finalidade).like('%ESTOQUE%'),1),else_=0)),
     func.count(func.distinct(func.nullif(Tag.localidade,''))),
     func.sum(case((db.or_(Tag.nf_entrada.is_(None),Tag.nf_entrada==''),1),else_=0)),
     func.sum(case((db.or_(Tag.patrimonio.is_(None),Tag.patrimonio==''),1),else_=0))
    ).first()
    product_filter=request.args.get('produto','').strip()
    issue_query=ImportIssue.query.filter(ImportIssue.resolved_at.is_(None))
    if product_filter:issue_query=issue_query.filter(ImportIssue.produto==product_filter)
    return jsonify(ok=True,total=total,page=page,limit=limit,items=[asdict(t) for t in items],groups=groups,
      summary={'total':total,'locacao':int(summary_row[0] or 0),'venda':int(summary_row[1] or 0),
      'estoque':int(summary_row[2] or 0),'localidades':int(summary_row[3] or 0),
      'sem_nf':int(summary_row[4] or 0),'sem_patrimonio':int(summary_row[5] or 0),
      'divergencias':issue_query.count()})
 @app.get('/api/engenharia/tags/<int:tag_id>/historico')
 @guard("engineering.tags.audit")
 def tags_history(tag_id):
    Tag.query.get_or_404(tag_id)
    rows=Audit.query.filter_by(tag_id=tag_id).order_by(Audit.created_at.desc(),Audit.id.desc()).limit(300).all()
    return jsonify(ok=True,items=[dict(id=a.id,usuario_id=a.user_id,acao=a.action,data=a.created_at.isoformat() if a.created_at else '',antes=json.loads(a.before_json) if a.before_json else None,depois=json.loads(a.after_json) if a.after_json else None) for a in rows])
 @app.post('/api/engenharia/tags/<int:tag_id>/editar')
 @guard("engineering.tags.edit")
 def tags_edit(tag_id):
    t=Tag.query.get_or_404(tag_id)
    payload=request.get_json(silent=True) or {}
    reason=str(payload.get('justificativa') or '').strip()
    if len(reason)<10:return jsonify(ok=False,error='Informe uma justificativa com pelo menos 10 caracteres.'),400
    changes=payload.get('alteracoes')
    if not isinstance(changes,dict):return jsonify(ok=False,error='Alterações inválidas.'),400
    allowed=set(FIELDS)
    if not changes or set(changes)-allowed:return jsonify(ok=False,error='Campos inválidos ou nenhuma alteração informada.'),400
    before=asdict(t);updated={}
    for field,value in changes.items():
     if not isinstance(value,(str,int,float)) and value is not None:return jsonify(ok=False,error='Valor inválido: '+field),400
     value=str(value).strip() if value is not None else ''
     if len(value)>(4000 if field=='observacao' else 200):return jsonify(ok=False,error='Campo muito longo: '+field),400
     if field in ('serial','produto') and not value:return jsonify(ok=False,error='Produto e TAG são obrigatórios.'),400
     if value!=str(getattr(t,field) or ''):updated[field]=value
    if not updated:return jsonify(ok=False,error='Nenhum valor foi alterado.'),400
    prod=updated.get('produto',t.produto);serial=updated.get('serial',t.serial)
    duplicate=Tag.query.filter(func.lower(Tag.produto)==prod.lower(),func.lower(Tag.serial)==serial.lower(),Tag.id!=t.id).first()
    if duplicate:return jsonify(ok=False,error='Já existe TAG com este produto e número de série.'),409
    try:
     for k,v in updated.items():setattr(t,k,v)
     t.updated_at=datetime.utcnow()
     after=asdict(t)
     audit=dict(campos=list(updated),justificativa=reason,registro=after)
     db.session.add(Audit(tag_id=t.id,user_id=int(session['user_id']),action='EDIT',before_json=json.dumps(before,ensure_ascii=False),after_json=json.dumps(audit,ensure_ascii=False)))
     db.session.commit()
     return jsonify(ok=True,tag=asdict(t),campos=list(updated))
    except Exception:
     db.session.rollback();app.logger.exception('Falha edição TAG');return jsonify(ok=False,error='Não foi possível salvar a alteração.'),500
 @app.get('/api/engenharia/tags/modelo')
 @guard("engineering.tags.import")
 def tags_model():
    wb=Workbook();ws=wb.active;ws.title='TAGs';ws.append(['Número de Série (TAG)','Patrimônio','Produto','Produção / Lote','Empresa','Local / Projeto','Finalidade','Status Operacional','NF Entrada','NF Saída','Caixa','Observação'])
    ws.append(['BF80-EXEMPLO-001','','BF80','1ª Produção','Autopass','BELÉM','Locação','Instalado','','','','EXEMPLO — excluir antes de importar'])
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='17345D');c.font=Font(color='FFFFFF',bold=True)
    ws.freeze_panes='A2';ws.auto_filter.ref='A1:L2'
    for col in 'ABCDEFGHIJKL':ws.column_dimensions[col].width=22
    b=BytesIO();wb.save(b);b.seek(0)
    return send_file(b,as_attachment=True,download_name='Modelo_TAGs_Produtos.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 @app.get('/api/engenharia/tags/exportar')
 @guard("engineering.tags.export")
 def tags_export():
    rows=query_filtered().order_by(Tag.produto,Tag.serial).all();wb=Workbook();ws=wb.active;ws.title='TAGs'
    ws.append(['Nº Série (TAG)','Patrimônio','Produto','Produção/Lote','Empresa','Local/Projeto','Finalidade','Status','NF Entrada','NF Saída','Caixa','Observação'])
    for t in rows:ws.append([getattr(t,k) or '' for k in FIELDS])
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='17345D');c.font=Font(color='FFFFFF',bold=True)
    ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:L{max(2,ws.max_row)}'
    for col in 'ABCDEFGHIJKL':ws.column_dimensions[col].width=23
    summary=wb.create_sheet('Resumo');summary.append(['Indicador','Quantidade']);summary.append(['TAGs',len(rows)])
    for k in ['produto','lote','localidade','finalidade']:
     summary.append([]);summary.append([k.upper(),'Quantidade']);counts={}
     for t in rows:counts[getattr(t,k) or 'Não informado']=counts.get(getattr(t,k) or 'Não informado',0)+1
     for name,n in sorted(counts.items(),key=lambda v:-v[1]):summary.append([name,n])
    summary.column_dimensions['A'].width=40;summary.column_dimensions['B'].width=18
    b=BytesIO();wb.save(b);b.seek(0)
    return send_file(b,as_attachment=True,download_name='Controle_TAGs_Produtos.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 def parse_xlsx(upload):
    if not upload or not upload.filename.lower().endswith('.xlsx'):raise ValueError('Envie um arquivo .xlsx válido.')
    raw=upload.read()
    if len(raw)>15*1024*1024:raise ValueError('Arquivo acima do limite de 15 MB.')
    wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
    try:
     best=None
     for ws in wb.worksheets:
      for ix,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,35),values_only=True),1):
       mapping={}
       for j,cell in enumerate(row):
        h=norm(cell)
        for key,alias in ALIASES.items():
         if h in alias and key not in mapping:mapping[key]=j
       if 'serial' in mapping and (best is None or len(mapping)>len(best[2])):best=(ws,ix,mapping)
     if best is None:raise ValueError('Coluna de número de série não identificada. Utilize o modelo padrão de importação.')
     ws,start,mapping=best
     rows=[];errors=[];seen={};duplicates={}
     for ix,row in enumerate(ws.iter_rows(min_row=start+1,values_only=True),start+1):
      def val(k):
       j=mapping.get(k);v=row[j] if j is not None and j<len(row) else None
       if v is None:return ''
       if isinstance(v,(int,float)) and float(v).is_integer():return str(int(v))
       return str(v).strip()
      item={k:val(k) for k in FIELDS};item['produto']=item['produto'] or request.form.get('produto','BF80').strip() or 'BF80'
      if not any(item.values()) or not item['serial']:continue
      key=(item['produto'].casefold(),item['serial'].casefold())
      if key in seen:
       duplicates.setdefault(key,{'produto':item['produto'],'serial':item['serial'],'linhas':[seen[key]],'detalhes':[]})['linhas'].append(ix)
      else:seen[key]=ix
      rows.append(item)
      if len(rows)>20000:raise ValueError('Limite de 20.000 TAGs por importação.')
     return rows,[{'produto':v['produto'],'serial':v['serial'],'linhas':v['linhas']} for v in duplicates.values()]
    finally:wb.close()
 @app.get('/api/engenharia/tags/divergencias')
 @guard("engineering.tags.view")
 def tags_issues():
    q=ImportIssue.query.filter(ImportIssue.resolved_at.is_(None))
    product=request.args.get('produto','').strip()
    if product:q=q.filter(ImportIssue.produto==product)
    rows=q.order_by(ImportIssue.created_at.desc()).limit(1000).all()
    return jsonify(ok=True,total=len(rows),items=[dict(id=x.id,produto=x.produto,serial=x.serial,linhas=x.lines,descricao=x.detail,detalhes_json=x.detail if x.detail.startswith('[') else '[]',arquivo=x.file_name,data=x.created_at.isoformat()) for x in rows])
 @app.post('/api/engenharia/tags/divergencias/<int:issue_id>/resolver')
 @guard("engineering.tags.issues")
 def tags_resolve_issue(issue_id):
    issue=ImportIssue.query.get_or_404(issue_id)
    if issue.resolved_at is not None:return jsonify(ok=False,error='Divergência já tratada.'),409
    payload=request.get_json(silent=True) or {}
    action=payload.get('acao')
    if action not in ('same','correct'):return jsonify(ok=False,error='Informe o tipo de tratamento.'),400
    uid=int(session['user_id'])
    if action=='correct':
     corrected=str(payload.get('serial_corrigido') or '').strip()
     if not corrected or corrected.casefold()==issue.serial.casefold():return jsonify(ok=False,error='A TAG corrigida deve ser diferente da original.'),400
     if len(corrected)>160:return jsonify(ok=False,error='TAG excede o tamanho permitido.'),400
     exists=Tag.query.filter(db.func.lower(Tag.produto)==issue.produto.lower(),db.func.lower(Tag.serial)==corrected.lower()).first()
     if exists:return jsonify(ok=False,error='TAG corrigida já cadastrada. Confira antes de prosseguir.'),409
     # Registro mínimo, pois as linhas conflitantes originais não foram armazenadas na REV25.
     tag=Tag(produto=issue.produto,serial=corrected)
     db.session.add(tag);db.session.flush()
     db.session.add(Audit(tag_id=tag.id,user_id=uid,action='INSERT',before_json=None,after_json=json.dumps({'produto':issue.produto,'serial':corrected,'origem':'tratamento de divergência','issue_id':issue.id},ensure_ascii=False)))
    issue.resolved_at=datetime.utcnow()
    db.session.commit()
    return jsonify(ok=True,acao=action,complementar_dados=(action=='correct'))
 @app.get('/api/engenharia/tags/<int:tag_id>/fotos')
 @guard('engineering.tags.view')
 def tags_photos_list(tag_id):
    if not db.session.get(Tag,tag_id):return jsonify(ok=False,error='TAG não encontrada'),404
    rows=TagPhoto.query.filter_by(tag_id=tag_id).order_by(TagPhoto.id.desc()).all()
    return jsonify(ok=True,items=[dict(id=x.id,descricao=x.description,usuario=x.user_id,data=x.created_at.isoformat() if x.created_at else '') for x in rows])
 @app.get('/api/engenharia/tags/<int:tag_id>/fotos/<int:photo_id>')
 @guard('engineering.tags.view')
 def tags_photo_view(tag_id,photo_id):
    item=TagPhoto.query.filter_by(id=photo_id,tag_id=tag_id).first_or_404()
    b=BytesIO(item.image);b.seek(0)
    response=send_file(b,mimetype=item.mime,as_attachment=False,download_name=f'tag_{tag_id}_foto_{photo_id}')
    response.headers['Cache-Control']='private, no-store'
    response.headers['X-Content-Type-Options']='nosniff'
    return response
 @app.post('/api/engenharia/tags/<int:tag_id>/fotos')
 @guard('engineering.tags.edit')
 def tags_photo_add(tag_id):
    tag=db.session.get(Tag,tag_id)
    if not tag:return jsonify(ok=False,error='TAG não encontrada'),404
    f=request.files.get('foto')
    if not f:return jsonify(ok=False,error='Selecione uma fotografia'),400
    data=f.read(3*1024*1024+1)
    if not data or len(data)>3*1024*1024:return jsonify(ok=False,error='Foto deve ter até 3 MB'),400
    try:
     with Image.open(BytesIO(data)) as im:
      im.verify()
     with Image.open(BytesIO(data)) as im:
      fmt=im.format
    except (UnidentifiedImageError,ValueError,OSError):return jsonify(ok=False,error='Arquivo de imagem inválido'),400
    mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp'}.get(fmt)
    if not mime:return jsonify(ok=False,error='Use JPG, PNG ou WebP'),400
    desc=(request.form.get('descricao') or '').strip()[:250]
    item=TagPhoto(tag_id=tag_id,image=data,mime=mime,description=desc,user_id=int(session['user_id']))
    db.session.add(item);db.session.flush()
    db.session.add(Audit(tag_id=tag_id,user_id=int(session['user_id']),action='PHOTO_ADD',after_json=json.dumps({'foto_id':item.id,'descricao':desc},ensure_ascii=False)))
    db.session.commit()
    return jsonify(ok=True,id=item.id),201
 @app.delete('/api/engenharia/tags/<int:tag_id>/fotos/<int:photo_id>')
 @guard('engineering.tags.edit')
 def tags_photo_delete(tag_id,photo_id):
    item=TagPhoto.query.filter_by(id=photo_id,tag_id=tag_id).first_or_404()
    db.session.add(Audit(tag_id=tag_id,user_id=int(session['user_id']),action='PHOTO_DELETE',before_json=json.dumps({'foto_id':item.id,'descricao':item.description},ensure_ascii=False)))
    db.session.delete(item);db.session.commit()
    return jsonify(ok=True)
 @app.get('/api/engenharia/tags/atualizacao-massa/modelo')
 @guard("engineering.tags.bulk")
 def tags_bulk_model():
    rows=query_filtered().order_by(Tag.produto,Tag.serial).all()
    wb=Workbook();ws=wb.active;ws.title='Atualizacao_TAGs'
    ws.append(['ID Interno','Produto','Nº Série (TAG)','Patrimônio','Produção / Lote','Empresa','Local / Projeto','Finalidade','Status Operacional','NF Entrada','NF Saída','Caixa','Observação'])
    cols=['id','produto','serial','patrimonio','lote','empresa','localidade','finalidade','status','nf_entrada','nf_saida','caixa','observacao']
    for t in rows:ws.append([getattr(t,k) or '' for k in cols])
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='17345D');c.font=Font(color='FFFFFF',bold=True)
    ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:M{max(2,ws.max_row)}'
    for col in 'ABCDEFGHIJKLM':ws.column_dimensions[col].width=24
    info=wb.create_sheet('Instrucoes');info.append(['Atualização em massa - Controle de TAGs']);info.append(['Não altere o ID Interno. Produto e TAG são usados para conferir a identidade.']);info.append(['Células vazias NÃO apagam informações.']);info.append(['Altere somente campos necessários. Justificativa será exigida na confirmação.']);info.append(['Para corrigir Produto ou Número de Série, utilize a edição individual auditada.']);info.column_dimensions['A'].width=115
    b=BytesIO();wb.save(b);b.seek(0)
    return send_file(b,as_attachment=True,download_name='Atualizacao_Massa_TAGs.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 @app.post('/api/engenharia/tags/atualizacao-massa')
 @guard("engineering.tags.bulk")
 def tags_bulk_update():
    upload=request.files.get('arquivo')
    if not upload or not upload.filename.lower().endswith('.xlsx'):return jsonify(ok=False,error='Envie um Excel .xlsx exportado para atualização em massa.'),400
    data=upload.read()
    if len(data)>15*1024*1024:return jsonify(ok=False,error='Arquivo acima de 15 MB.'),400
    try:wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    except Exception:return jsonify(ok=False,error='Não foi possível ler a planilha.'),400
    cols=['id','produto','serial','patrimonio','lote','empresa','localidade','finalidade','status','nf_entrada','nf_saida','caixa','observacao']
    try:
     ws=wb['Atualizacao_TAGs'] if 'Atualizacao_TAGs' in wb.sheetnames else wb.active
     headers=[norm(v) for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
     # Compatibilidade com o cabeçalho gerado pelo exportador e variações de Excel.
     headers=[('N_SERIE_TAG' if h in ('N_SERIE_TAG','NO_SERIE_TAG','NUMERO_SERIE_TAG','N_DE_SERIE_TAG') else h) for h in headers]
     if headers[:3]!=['ID_INTERNO','PRODUTO','N_SERIE_TAG']:
      return jsonify(ok=False,error='Cabeçalhos incompatíveis. Esperado: ID Interno, Produto, Nº Série (TAG). Encontrado: '+', '.join(headers[:3])),400
     expected=['ID_INTERNO','PRODUTO','N_SERIE_TAG','PATRIMONIO','PRODUCAO_LOTE','EMPRESA','LOCAL_PROJETO','FINALIDADE','STATUS_OPERACIONAL','NF_ENTRADA','NF_SAIDA','CAIXA','OBSERVACAO']
     if any(h not in headers for h in expected):
      missing=[h for h in expected if h not in headers]
      return jsonify(ok=False,error='Colunas ausentes ou renomeadas: '+', '.join(missing)),400
     mapping={k:headers.index(expected[i]) for i,k in enumerate(cols)}
     changes=[];issues=[];seen=set();examined=0;unchanged=0
     for ix,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
      if not any(v is not None and str(v).strip() for v in row):continue
      examined+=1
      if examined>20000:return jsonify(ok=False,error='Limite de 20.000 linhas.'),400
      try:ident=int(row[mapping['id']]);assert ident>0
      except Exception:issues.append({'linha':ix,'motivo':'ID interno inválido'});continue
      if ident in seen:issues.append({'linha':ix,'motivo':'ID duplicado na planilha'});continue
      seen.add(ident);t=db.session.get(Tag,ident)
      if not t:issues.append({'linha':ix,'motivo':'TAG não encontrada no banco'});continue
      def clean(v):
       if v is None:return ''
       if isinstance(v,(int,float)) and float(v).is_integer():return str(int(v))
       return str(v).strip()
      product=clean(row[mapping['produto']] if len(row)>mapping['produto'] else None);serial=clean(row[mapping['serial']] if len(row)>mapping['serial'] else None)
      if product.casefold()!=t.produto.casefold() or serial.casefold()!=t.serial.casefold():
       issues.append({'linha':ix,'motivo':'Produto ou série diferente do registro de ID '+str(ident)});continue
      modified={}
      for k,j in mapping.items():
       if k in ('id','produto','serial') or j>=len(row):continue
       v=clean(row[j])
       if v and v!=str(getattr(t,k) or ''):
        if len(v)>(2000 if k=='observacao' else 200):issues.append({'linha':ix,'motivo':'Valor excessivamente longo em '+k});modified={};break
        modified[k]=v
      if modified:changes.append({'id':ident,'produto':t.produto,'serial':t.serial,'linha':ix,'antes':{k:str(getattr(t,k) or '') for k in modified},'depois':modified,'updated_at':t.updated_at.isoformat() if t.updated_at else ''})
      else:unchanged+=1
    finally:wb.close()
    result=dict(ok=True,preview=True,analisadas=examined,alteradas=len(changes),sem_alteracao=unchanged,divergencias=len(issues),problemas=issues[:200],alteracoes=changes[:200])
    if request.form.get('confirmar')!='sim':return jsonify(**result)
    reason=(request.form.get('justificativa') or '').strip()
    if len(reason)<10:return jsonify(ok=False,error='Justificativa obrigatória com pelo menos 10 caracteres.'),400
    if issues:return jsonify(ok=False,error='Existem divergências. Corrija a planilha e valide novamente antes de gravar.',problemas=issues[:200]),409
    if not changes:return jsonify(ok=False,error='Nenhuma alteração identificada.'),400
    try:
     for change in changes:
      t=db.session.get(Tag,change['id'])
      if not t or (t.updated_at.isoformat() if t.updated_at else '')!=change['updated_at']:
       db.session.rollback();return jsonify(ok=False,error='Dados alterados durante a importação. Reexporte a planilha e valide novamente.'),409
      before=asdict(t)
      for k,v in change['depois'].items():setattr(t,k,v)
      t.updated_at=datetime.utcnow()
      db.session.add(Audit(tag_id=t.id,user_id=int(session['user_id']),action='BULK_UPDATE',before_json=json.dumps(before,ensure_ascii=False),after_json=json.dumps({'registro':asdict(t),'campos':list(change['depois']),'justificativa':reason,'arquivo':upload.filename[:255],'linha':change['linha']},ensure_ascii=False)))
     db.session.commit()
    except Exception:
     db.session.rollback();app.logger.exception('Erro na atualização em massa de TAGs');return jsonify(ok=False,error='Falha na gravação. Transação revertida.'),500
    return jsonify(ok=True,preview=False,atualizadas=len(changes),justificativa=reason)
 @app.post('/api/engenharia/tags/importar')
 @guard("engineering.tags.import")
 def tags_import():
    try:rows,errors=parse_xlsx(request.files.get('arquivo'))
    except Exception as e:return jsonify(ok=False,error=str(e)),400
    if not rows:return jsonify(ok=False,error='Nenhuma TAG encontrada.'),400
    duplicate_keys={(d['produto'].casefold(),d['serial'].casefold()) for d in errors}
    valid=[x for x in rows if (x['produto'].casefold(),x['serial'].casefold()) not in duplicate_keys]
    existing={ (t.produto.casefold(),t.serial.casefold()):t for t in Tag.query.filter(Tag.produto.in_({x['produto'] for x in valid})).all() }
    inserts=sum((x['produto'].casefold(),x['serial'].casefold()) not in existing for x in valid)
    updates=len(valid)-inserts
    if updates and not can('engineering.tags.bulk'):
     return jsonify(ok=False,error='A planilha contém TAGs existentes. Atualização requer permissão específica e o fluxo de Atualização em massa.'),403
    if updates:
     return jsonify(ok=False,error='Para atualizar TAGs existentes, utilize Atualização em massa com justificativa e auditoria.'),409
    preview=dict(ok=True,preview=True,inserts=inserts,updates=updates,total=len(rows),validos=len(valid),divergencias=errors,divergencias_count=len(errors),ignorados=len(rows)-len(valid))
    if request.form.get('confirmar')!='sim':return jsonify(**preview)
    if errors and request.form.get('ignorar_divergencias')!='sim':
     return jsonify(ok=False,error='Confirme explicitamente a importação parcial: TAGs duplicadas serão isoladas para conferência.'),409
    uid=int(session['user_id'])
    try:
     for item in valid:
      key=(item['produto'].casefold(),item['serial'].casefold());t=existing.get(key)
      before=json.dumps(asdict(t),ensure_ascii=False) if t else None
      if not t:
       t=Tag(serial=item['serial'],produto=item['produto']);db.session.add(t);db.session.flush();existing[key]=t
      for k in FIELDS:
       if k in ('serial','produto'):continue
       # Em atualização, células vazias não apagam informações existentes.
       if item[k]:setattr(t,k,item[k])
      db.session.add(Audit(tag_id=t.id,user_id=uid,action='UPDATE' if before else 'INSERT',before_json=before,after_json=json.dumps(asdict(t),ensure_ascii=False)))
     for issue in errors:
      # Preserva histórico; repetição da mesma planilha não cria alertas duplicados.
      found=ImportIssue.query.filter_by(produto=issue['produto'],serial=issue['serial'],file_name=(request.files.get('arquivo').filename or '')[:255],resolved_at=None).first()
      if not found:db.session.add(ImportIssue(produto=issue['produto'],serial=issue['serial'],lines=', '.join(map(str,issue['linhas'])),detail='TAG repetida dentro do Excel; não importada automaticamente',file_name=(request.files.get('arquivo').filename or '')[:255],user_id=uid))
     db.session.commit()
    except Exception:
     db.session.rollback();app.logger.exception('Erro importação TAGs');return jsonify(ok=False,error='Falha ao gravar. Nenhuma alteração foi confirmada.'),500
    return jsonify(ok=True,inserted=inserts,updated=updates,total=len(valid),divergencias_count=len(errors),ignorados=len(rows)-len(valid))


