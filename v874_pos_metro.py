"""V87.4 incremental: atividade POS Metrô. Não altera tabelas legadas."""
from datetime import datetime, timedelta
from io import BytesIO
from functools import wraps
from flask import request,session,abort,redirect,jsonify,render_template,send_file
from sqlalchemy import func, or_, text
import re
from openpyxl import load_workbook,Workbook

def install_pos_metro(app,db,can):
 class POS(db.Model):
  __tablename__='v874_pos_metro'
  id=db.Column(db.Integer,primary_key=True)
  terminal=db.Column(db.String(90),unique=True,nullable=False)
  line=db.Column(db.String(100),nullable=False,default='')
  station=db.Column(db.String(150),nullable=False,default='')
  serial=db.Column(db.String(120))
  asset=db.Column(db.String(120))
  manufacturer=db.Column(db.String(120))
  model=db.Column(db.String(120))
  status=db.Column(db.String(25),nullable=False,default='Pendente')
  installed_on=db.Column(db.DateTime)
  terminal_number=db.Column(db.String(40))
 class Event(db.Model):
  __tablename__='v874_pos_metro_events'
  id=db.Column(db.Integer,primary_key=True)
  pos_id=db.Column(db.Integer,db.ForeignKey('v874_pos_metro.id'),nullable=False,index=True)
  old_serial=db.Column(db.String(120))
  new_serial=db.Column(db.String(120))
  old_model=db.Column(db.String(120))
  new_model=db.Column(db.String(120))
  notes=db.Column(db.Text)
  user_id=db.Column(db.Integer,nullable=False)
  occurred_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
 class Photo(db.Model):
  __tablename__='v876_pos_metro_photos'
  id=db.Column(db.Integer,primary_key=True)
  event_id=db.Column(db.Integer,db.ForeignKey('v874_pos_metro_events.id'),nullable=False,index=True)
  filename=db.Column(db.String(160),nullable=False)
  mimetype=db.Column(db.String(40),nullable=False)
  content=db.Column(db.LargeBinary,nullable=False)
  created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
 class AP(db.Model):
  __tablename__='v8738_pos_metro_ap'
  id=db.Column(db.Integer,primary_key=True)
  line=db.Column(db.String(100),nullable=False)
  station=db.Column(db.String(150),nullable=False)
  manufacturer=db.Column(db.String(50),nullable=False,default='Mercusys')
  model=db.Column(db.String(50),nullable=False,default='MW301R')
  status=db.Column(db.String(25),nullable=False,default='Pendente')
  installed_on=db.Column(db.DateTime)
  updated_by=db.Column(db.Integer)
  notes=db.Column(db.Text)
  __table_args__=(db.UniqueConstraint('line','station',name='uq_v8738_ap_station'),)
 class APEvent(db.Model):
  __tablename__='v8738_pos_metro_ap_events'
  id=db.Column(db.Integer,primary_key=True)
  ap_id=db.Column(db.Integer,db.ForeignKey('v8738_pos_metro_ap.id'),nullable=False,index=True)
  status=db.Column(db.String(25),nullable=False)
  notes=db.Column(db.Text)
  user_id=db.Column(db.Integer,nullable=False)
  occurred_at=db.Column(db.DateTime,nullable=False,default=datetime.utcnow)
 class APPhoto(db.Model):
  __tablename__='v8738_pos_metro_ap_photos'
  id=db.Column(db.Integer,primary_key=True)
  ap_id=db.Column(db.Integer,db.ForeignKey('v8738_pos_metro_ap.id'),nullable=False,index=True)
  filename=db.Column(db.String(160),nullable=False)
  mimetype=db.Column(db.String(40),nullable=False)
  content=db.Column(db.LargeBinary,nullable=False)
  user_id=db.Column(db.Integer,nullable=False)
  created_at=db.Column(db.DateTime,nullable=False,default=datetime.utcnow)
 def ensure_schema():
  POS.__table__.create(bind=db.engine,checkfirst=True)
  Event.__table__.create(bind=db.engine,checkfirst=True)
  Photo.__table__.create(bind=db.engine,checkfirst=True)
  AP.__table__.create(bind=db.engine,checkfirst=True)
  APPhoto.__table__.create(bind=db.engine,checkfirst=True)
  APEvent.__table__.create(bind=db.engine,checkfirst=True)
  if db.engine.dialect.name=='postgresql':
   with db.engine.begin() as conn:
    conn.execute(text('ALTER TABLE v874_pos_metro ADD COLUMN IF NOT EXISTS terminal_number VARCHAR(40)'))
  elif db.engine.dialect.name=='sqlite':
   with db.engine.begin() as conn:
    columns=[r[1] for r in conn.execute(text('PRAGMA table_info(v874_pos_metro)'))]
    if 'terminal_number' not in columns:conn.execute(text('ALTER TABLE v874_pos_metro ADD COLUMN terminal_number VARCHAR(40)'))
 def short_number(p):
  return p.terminal_number or (re.search(r'(\d+)$',p.terminal).group(1) if re.search(r'(\d+)$',p.terminal) else p.terminal)
 def guard(permission):
  def deco(fn):
   @wraps(fn)
   def wrapped(*args,**kwargs):
    if not session.get('user_id'):return redirect('/login')
    if not can(permission):abort(403)
    ensure_schema()
    return fn(*args,**kwargs)
   return wrapped
  return deco
 def date_range():
  start=(request.args.get('data_inicio') or '').strip();end=(request.args.get('data_fim') or '').strip()
  try:
   a=datetime.strptime(start,'%Y-%m-%d') if start else None
   b=datetime.strptime(end,'%Y-%m-%d')+timedelta(days=1) if end else None
   if a and b and a>=b:abort(400,description='Data inicial posterior à data final')
   return a,b
  except ValueError:abort(400,description='Formato de data inválido')
 def date_filtered(q):
  a,b=date_range()
  if a or b:
   # Date of actual completion; avoid using the asset registration date.
   q=q.filter(POS.status=='Concluído',POS.installed_on.isnot(None))
   if a:q=q.filter(POS.installed_on>=a)
   if b:q=q.filter(POS.installed_on<b)
  return q
 def selected():
  q=POS.query
  for arg,col in [('linha',POS.line),('estacao',POS.station),('status',POS.status)]:
   v=(request.args.get(arg) or '').strip()
   if v:q=q.filter(col==v)
  search=(request.args.get('q') or '').strip()
  if search:q=q.filter(or_(POS.terminal.ilike('%'+search+'%'),POS.terminal_number.ilike('%'+search+'%'),POS.serial.ilike('%'+search+'%'),POS.asset.ilike('%'+search+'%')))
  return date_filtered(q)
 @app.get('/field/troca-pos-metro')
 @guard('field.pos_metro')
 def pos_page():
  q=selected();items=q.order_by(POS.line,POS.station,POS.terminal).limit(500).all()
  overall=POS.query
  total=overall.count();done=overall.filter(POS.status=='Concluído').count();ongoing=overall.filter(POS.status=='Em andamento').count()
  lines=[x[0] for x in db.session.query(POS.line).distinct().order_by(POS.line)]
  stations_q=db.session.query(POS.station)
  if request.args.get('linha'):stations_q=stations_q.filter(POS.line==request.args['linha'])
  stations=[x[0] for x in stations_q.distinct().order_by(POS.station)]
  station_name=(request.args.get('estacao') or '').strip()
  station_q=POS.query.filter(POS.station==station_name) if station_name else None
  if station_q is not None and request.args.get('linha'):station_q=station_q.filter(POS.line==request.args['linha'])
  st_total=station_q.count() if station_q is not None else 0
  st_done=station_q.filter(POS.status=='Concluído').count() if station_q is not None else 0
  st_ongoing=station_q.filter(POS.status=='Em andamento').count() if station_q is not None else 0
  known_models=[x[0] for x in db.session.query(POS.model).filter(POS.model.isnot(None),POS.model!='').distinct().order_by(POS.model).limit(300)]
  known_manufacturers=[x[0] for x in db.session.query(POS.manufacturer).filter(POS.manufacturer.isnot(None),POS.manufacturer!='').distinct().order_by(POS.manufacturer).limit(300)]
  metric_rows=[{'id':x.id,'line':x.line,'station':x.station,'terminal':x.terminal,'terminal_number':short_number(x),'model':x.model or '', 'manufacturer':x.manufacturer or '', 'serial':x.serial or '', 'asset':x.asset or '', 'status':x.status,'installed_on':x.installed_on.strftime('%Y-%m-%d') if x.installed_on else ''} for x in POS.query.order_by(POS.line,POS.station,POS.terminal).all()]
  ap_station=AP.query.filter_by(line=request.args.get('linha'),station=station_name).first() if station_name and request.args.get('linha') else None
  if station_name and request.args.get('linha') and st_total:
   st_total+=1
   if ap_station and ap_station.status=='Concluído':st_done+=1
   if ap_station and ap_station.status=='Em andamento':st_ongoing+=1
  return render_template('v874_pos_metro.html',ap_station=ap_station,metric_rows=metric_rows,known_models=known_models,known_manufacturers=known_manufacturers,items=items,short_number=short_number,total=total,done=done,ongoing=ongoing,lines=lines,stations=stations,can_pos=can,st_total=st_total,st_done=st_done,st_ongoing=st_ongoing,filtered_count=q.count())
 @app.post('/field/troca-pos-metro/importar')
 @guard('field.pos_metro.import')
 def pos_import():
  f=request.files.get('arquivo')
  if not f or not (f.filename or '').lower().endswith('.xlsx'):return jsonify(error='Arquivo XLSX necessário'),400
  raw=f.read(5_000_001)
  if len(raw)>5_000_000:return jsonify(error='Arquivo excede 5 MB'),413
  try:wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
  except Exception:return jsonify(error='Planilha inválida'),400
  if 'Metrô' not in wb and 'POS Metrô' not in wb:return jsonify(error='Aba Metrô ou POS Metrô não encontrada'),400
  sheet=wb['Metrô'] if 'Metrô' in wb else wb['POS Metrô']; rows=sheet.iter_rows(values_only=True)
  header=next(rows,None)
  if not header:return jsonify(error='Planilha vazia'),400
  def norm(v):
   import unicodedata
   return ''.join(c for c in unicodedata.normalize('NFKD',str(v or '').strip().upper()) if not unicodedata.combining(c)).replace('º','').replace('°','').replace('°','').replace(' ','').replace('N','N')
  columns={norm(v):i for i,v in enumerate(header)}
  required=['TERMINAL','LINHA','ESTACAO']
  if any(k not in columns for k in required):return jsonify(error='Colunas obrigatórias ausentes: TERMINAL, LINHA, ESTAÇÃO'),400
  def val(r,k):
   i=columns.get(k)
   return str(r[i]).strip() if i is not None and i<len(r) and r[i] is not None else ''
  records={}
  for r in rows:
   terminal=val(r,'TERMINAL')[:90];line=val(r,'LINHA')[:100];station=val(r,'ESTACAO')[:150]
   if not terminal or not line or not station:continue
   if terminal in records:return jsonify(error='Terminal duplicado na planilha: '+terminal),400
   number=(val(r,'NTERMINAL2') or val(r,'NUMEROTERMINAL') or val(r,'NTERMINAL') or (re.search(r'(\d+)$',terminal).group(1) if re.search(r'(\d+)$',terminal) else ''))[:40]
   records[terminal]=(line,station,number,val(r,'MODELO')[:120],(val(r,'NSERIE') or val(r,'NoSERIE') or val(r,'NUMERODESERIE'))[:120],val(r,'PATRIMONIO')[:120],val(r,'FABRICANTE')[:120])
  if not records:return jsonify(error='Nenhum POS válido encontrado'),400
  try:
   existing={p.terminal:p for p in POS.query.filter(POS.terminal.in_(list(records))).all()}
   for terminal,(line,station,number,model,serial,asset,manufacturer) in records.items():
    obj=existing.get(terminal)
    if obj is None:
     obj=POS(terminal=terminal,status='Pendente');db.session.add(obj)
    obj.line=line;obj.station=station;obj.terminal_number=number
    # Importações subsequentes não sobrescrevem informações executadas em campo.
    if obj.id is None or not db.session.query(Event.id).filter(Event.pos_id==obj.id).first():
     if model:obj.model=model
     if serial:obj.serial=serial
     if asset:obj.asset=asset
     if manufacturer:obj.manufacturer=manufacturer
   db.session.commit()
  except Exception:
   db.session.rollback();raise
  finally:wb.close()
  from flask import flash
  flash(f'Importação concluída: {len(records)} terminais processados.', 'success')
  return redirect('/field/troca-pos-metro')
 @app.post('/field/troca-pos-metro/<int:pid>/registrar')
 @guard('field.pos_metro.execute')
 def pos_register(pid):
  obj=db.session.get(POS,pid)
  if not obj:abort(404)
  status=(request.form.get('status') or 'Pendente').strip()
  if status not in ('Pendente','Em andamento','Concluído'):return jsonify(error='Status inválido'),400
  serial=(request.form.get('serial') or '').strip()[:120]
  # Série e patrimônio não são exigidos para registrar a instalação.
  manufacturer=(request.form.get('fabricante') or '').strip()
  models={'Gertec':'G820','Ingenico':'DX8000'}
  if manufacturer not in models:return jsonify(error='Selecione Gertec ou Ingenico'),400
  model=models[manufacturer]
  asset=obj.asset or ''
  photos=request.files.getlist('fotos')
  if len(photos)>5:return jsonify(error='Máximo de 5 fotos'),400
  valid=[]
  for photo in photos:
   if not photo.filename:continue
   data=photo.read(3_000_001)
   if len(data)>3_000_000:return jsonify(error='Cada foto deve ter até 3 MB'),400
   from PIL import Image
   try:
    im=Image.open(BytesIO(data));im.verify()
    if im.format not in ('JPEG','PNG','WEBP'):raise ValueError('Formato inválido')
   except Exception:return jsonify(error='Imagem inválida; utilize JPEG, PNG ou WEBP'),400
   mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp'}[im.format]
   valid.append((photo.filename[:160],mime,data))
  if status=='Concluído' and not valid and not db.session.query(Photo.id).join(Event,Photo.event_id==Event.id).filter(Event.pos_id==pid).first():return jsonify(error='Foto obrigatória para concluir POS'),400
  try:
   event=Event(pos_id=pid,old_serial=obj.serial,old_model=obj.model,new_serial=serial or obj.serial,new_model=model or obj.model,notes=(request.form.get('observacoes') or '')[:2000],user_id=session['user_id'])
   db.session.add(event);db.session.flush()
   for filename,mime,data in valid:db.session.add(Photo(event_id=event.id,filename=filename,mimetype=mime,content=data))
   if serial:obj.serial=serial
   if model:obj.model=model
   if manufacturer:obj.manufacturer=manufacturer
   if asset:obj.asset=asset
   obj.status=status
   if status=='Concluído':obj.installed_on=datetime.utcnow()
   db.session.commit()
  except Exception:
   db.session.rollback();raise
  return redirect('/field/troca-pos-metro')
 def validate_photos(files):
  from PIL import Image
  result=[]
  if len(files)>5:abort(400,description='Máximo de 5 fotos')
  for f in files:
   if not f.filename:continue
   data=f.read(3_000_001)
   if len(data)>3_000_000:abort(400,description='Foto acima de 3 MB')
   try:
    im=Image.open(BytesIO(data));im.verify()
    if im.format not in ('JPEG','PNG','WEBP'):raise ValueError()
    mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp'}[im.format]
   except Exception:abort(400,description='Foto inválida; use JPEG, PNG ou WEBP')
   result.append((f.filename[:160],mime,data))
  return result
 @app.post('/field/troca-pos-metro/novo-pos')
 @guard('field.pos_metro.execute')
 def pos_new():
  line=(request.form.get('linha') or '').strip()[:100]
  station=(request.form.get('estacao') or '').strip()[:150]
  terminal=(request.form.get('terminal') or '').strip()[:90]
  if not all([line,station,terminal]):return jsonify(error='Linha, estação e identificação do POS são obrigatórios'),400
  if not POS.query.filter_by(line=line,station=station).first():return jsonify(error='Estação não encontrada na base POS'),400
  if POS.query.filter(func.lower(POS.terminal)==terminal.lower()).first():return jsonify(error='POS já cadastrado na base, inclusive em outra estação'),409
  obj=POS(line=line,station=station,terminal=terminal,terminal_number=(request.form.get('numero') or '').strip()[:40],status='Pendente')
  db.session.add(obj);db.session.flush()
  db.session.add(Event(pos_id=obj.id,notes='Cadastro provisório em campo; pendente de validação do gestor',user_id=session['user_id']))
  db.session.commit()
  return redirect('/field/troca-pos-metro?linha='+__import__('urllib.parse',fromlist=['quote']).quote(line)+'&estacao='+__import__('urllib.parse',fromlist=['quote']).quote(station))
 @app.post('/field/troca-pos-metro/ap/registrar')
 @guard('field.pos_metro.execute')
 def ap_register():
  line=(request.form.get('linha') or '').strip()
  station=(request.form.get('estacao') or '').strip()
  if not line or not station or not POS.query.filter_by(line=line,station=station).first():return jsonify(error='Selecione uma estação válida'),400
  status=(request.form.get('status') or 'Pendente').strip()
  if status not in ('Pendente','Em andamento','Concluído'):return jsonify(error='Status inválido'),400
  valid=validate_photos(request.files.getlist('fotos'))
  ap=AP.query.filter_by(line=line,station=station).first()
  if status=='Concluído' and not valid and not (ap and APPhoto.query.filter_by(ap_id=ap.id).first()):return jsonify(error='Foto obrigatória para concluir o AP'),400
  if ap is None:
   ap=AP(line=line,station=station);db.session.add(ap);db.session.flush()
  ap.status=status;ap.updated_by=session['user_id'];ap.notes=(request.form.get('observacoes') or '')[:2000]
  db.session.add(APEvent(ap_id=ap.id,status=status,notes=ap.notes,user_id=session['user_id']))
  if status=='Concluído':ap.installed_on=datetime.utcnow()
  for name,mime,data in valid:db.session.add(APPhoto(ap_id=ap.id,filename=name,mimetype=mime,content=data,user_id=session['user_id']))
  db.session.commit()
  from urllib.parse import urlencode
  return redirect('/field/troca-pos-metro?'+urlencode({'linha':line,'estacao':station}))
 @app.get('/field/troca-pos-metro/ap/foto/<int:photo_id>')
 @guard('field.pos_metro')
 def ap_photo(photo_id):
  photo=db.session.get(APPhoto,photo_id)
  if not photo:abort(404)
  return send_file(BytesIO(photo.content),mimetype=photo.mimetype,download_name=photo.filename)
 @app.get('/field/troca-pos-metro/foto/<int:photo_id>')
 @guard('field.pos_metro')
 def pos_photo(photo_id):
  photo=db.session.get(Photo,photo_id)
  if not photo:abort(404)
  return send_file(BytesIO(photo.content),mimetype=photo.mimetype,download_name=photo.filename)
 @app.get('/field/troca-pos-metro/modelo')
 @guard('field.pos_metro.export')
 def pos_model():
  return build_export(empty=True)
 def build_export(empty=False):
  wb=Workbook();ws=wb.active;ws.title='POS Metrô'
  ws.append(['Cliente','Linha','Estação','N° TERMINAL2','Terminal','Modelo','Fabricante','Nº Série','Patrimônio','Status','Técnico responsável','Data da atividade','Observações'])
  if not empty:
   for x in selected().order_by(POS.line,POS.station,POS.terminal):
    event=Event.query.filter_by(pos_id=x.id).order_by(Event.occurred_at.desc()).first()
    ws.append(['METRÔ',x.line,x.station,short_number(x),x.terminal,x.model,x.manufacturer,x.serial,x.asset,x.status,event.user_id if event else '',event.occurred_at if event else '',event.notes if event else ''])
  stream=BytesIO();wb.save(stream);stream.seek(0)
  return send_file(stream,as_attachment=True,download_name=('modelo_pos_metro.xlsx' if empty else 'pos_metro_filtrado.xlsx'),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 @app.get('/field/troca-pos-metro/exportar')
 @guard('field.pos_metro.export')
 def pos_export():
  return build_export()
